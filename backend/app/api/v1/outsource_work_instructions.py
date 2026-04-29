from __future__ import annotations

from datetime import date, datetime, timezone
from functools import lru_cache
from io import BytesIO
from pathlib import Path
from typing import List
from uuid import uuid4

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile, status as http_status
from fastapi.responses import StreamingResponse
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from sqlalchemy import exists, func, or_, select
from sqlalchemy.orm import Session, aliased

from app.core.config import settings
from app.db.session import get_db
from app.models.lot import Lot
from app.models.order_line import OrderLine
from app.models.outsource_purchase_order import OutsourcePurchaseOrder
from app.models.outsource_purchase_order_item import OutsourcePurchaseOrderItem
from app.models.outsource_work_group import OutsourceWorkGroup
from app.models.outsource_work_group_item import OutsourceWorkGroupItem
from app.models.outsource_work_instruction import OutsourceWorkInstruction
from app.models.outsource_work_instruction_file import OutsourceWorkInstructionFile
from app.models.outsource_work_instruction_item import OutsourceWorkInstructionItem
from app.models.partner import Partner
from app.models.product import Product
from app.models.routing_template import RoutingTemplate
from app.schemas.outsource_work_instruction import (
    BohyunOutsourceGroupItemOut,
    BohyunOutsourceGroupListItemOut,
    BohyunOutsourceGroupListOut,
    BohyunOutsourceGroupShipBatch,
    BohyunOutsourceGroupWorkDone,
    OutsourcePurchaseOrderCreate,
    OutsourcePurchaseOrderCutSnapshot,
    OutsourcePurchaseOrderItemOut,
    OutsourcePurchaseOrderListItemOut,
    OutsourcePurchaseOrderListOut,
    OutsourcePurchaseOrderOut,
    OutsourcePurchaseOrderPrintSnapshot,
    OutsourcePurchaseOrderTargetListOut,
    OutsourcePurchaseOrderTargetOut,
    OutsourcePurchaseOrderWorkDone,
    OutsourceWorkInstructionBatchCreate,
    OutsourceWorkInstructionBatchOut,
    OutsourceWorkInstructionCandidateLotListOut,
    OutsourceWorkInstructionCandidateLotOut,
    OutsourceWorkInstructionCreate,
    OutsourceWorkInstructionFileOut,
    OutsourceWorkInstructionGroupCreate,
    OutsourceWorkInstructionGroupItemCreate,
    OutsourceWorkInstructionItemOut,
    OutsourceWorkInstructionOut,
    OutsourceWorkInstructionPlateUploadOut,
)

router = APIRouter(prefix="/outsource-work-instructions", tags=["OutsourceWorkInstruction"])

BOHYUN_DB_STATUS_VENDOR_RECEIVED = "VENDOR_RECEIVED"
BOHYUN_DB_STATUS_WORK_DONE = "WORK_DONE"
BOHYUN_DB_STATUS_SHIPPED = "SHIPPED"

BOHYUN_UI_STATUS_WAITING_INBOUND = "WAITING_INBOUND"
BOHYUN_UI_STATUS_INBOUNDED = "INBOUNDED"
BOHYUN_UI_STATUS_WORK_DONE = "WORK_DONE"
BOHYUN_UI_STATUS_SHIPPED = "SHIPPED"


def _to_bohyun_ui_status(db_status: str | None) -> str:
    if db_status == BOHYUN_DB_STATUS_VENDOR_RECEIVED:
        return BOHYUN_UI_STATUS_INBOUNDED

    if db_status == BOHYUN_DB_STATUS_WORK_DONE:
        return BOHYUN_UI_STATUS_WORK_DONE

    if db_status == BOHYUN_DB_STATUS_SHIPPED:
        return BOHYUN_UI_STATUS_SHIPPED

    return BOHYUN_UI_STATUS_WAITING_INBOUND


def _get_outsource_partner_name_by_process_type(process_type: str) -> str | None:
    normalized = (process_type or "").strip().upper()

    if normalized == "CUT":
        return "코리아라벨 주식회사"

    if normalized == "PRINT":
        return "주식회사 상림크리에이티브"

    if normalized == "DIECUT":
        return "보현문화"

    return None


def _get_bohyun_inbound_source_name(process_type: str) -> str | None:
    return _get_outsource_partner_name_by_process_type(process_type)


def _get_outsource_partner_by_process_type(
    db: Session,
    process_type: str,
) -> Partner | None:
    partner_name = _get_outsource_partner_name_by_process_type(process_type)

    if not partner_name:
        return None

    return (
        db.execute(
            select(Partner)
            .where(Partner.name == partner_name)
            .where(Partner.partner_type == "VENDOR")
            .where(Partner.is_active.is_(True))
        )
        .scalar_one_or_none()
    )


def _require_outsource_partner_by_process_type(
    db: Session,
    process_type: str,
) -> Partner:
    outsource_partner = _get_outsource_partner_by_process_type(
        db=db,
        process_type=process_type,
    )

    if outsource_partner is None:
        partner_name = _get_outsource_partner_name_by_process_type(process_type)

        raise HTTPException(
            status_code=409,
            detail=f"외주처 정보가 없습니다. partner 테이블에 [{partner_name}] VENDOR 거래처를 등록해주세요.",
        )

    return outsource_partner


def _normalize_ext(filename: str) -> str:
    return Path(filename).suffix.lower().strip()


def _sanitize_filename(filename: str) -> str:
    name = Path(filename).name
    return name.replace(" ", "_")


def _validate_plate_data_upload(file: UploadFile, size_bytes: int) -> None:
    original_name = file.filename or ""
    ext = _normalize_ext(original_name)

    if not original_name.strip():
        raise HTTPException(status_code=400, detail="File name is required")

    allowed_ext = settings.PLATE_DATA_ALLOWED_EXT
    if allowed_ext and ext not in allowed_ext:
        raise HTTPException(
            status_code=409,
            detail=f"File extension not allowed: {ext}",
        )

    max_bytes = settings.PLATE_DATA_MAX_MB * 1024 * 1024
    if size_bytes > max_bytes:
        raise HTTPException(
            status_code=409,
            detail=f"File size exceeded: max {settings.PLATE_DATA_MAX_MB}MB",
        )


def _build_plate_data_path(filename: str, uploaded_at: datetime) -> Path:
    safe_name = _sanitize_filename(filename)
    y = f"{uploaded_at.year:04d}"
    m = f"{uploaded_at.month:02d}"
    d = f"{uploaded_at.day:02d}"

    root = Path(settings.PLATE_DATA_STORAGE_ROOT)
    folder = root / "plate_data" / y / m / d
    folder.mkdir(parents=True, exist_ok=True)

    stored_name = f"{uuid4().hex}_{safe_name}"
    return folder / stored_name


def _generate_instruction_no(db: Session, instruction_date: date) -> str:
    yy = f"{instruction_date.year % 100:02d}"
    mm = f"{instruction_date.month:02d}"
    dd = f"{instruction_date.day:02d}"
    prefix = f"OWI{yy}{mm}{dd}"

    last = (
        db.execute(
            select(OutsourceWorkInstruction.instruction_no)
            .where(OutsourceWorkInstruction.instruction_no.like(f"{prefix}%"))
            .order_by(OutsourceWorkInstruction.instruction_no.desc())
            .limit(1)
        )
        .scalar_one_or_none()
    )

    if not last:
        seq = 1
    else:
        try:
            seq = int(last[-3:]) + 1
        except ValueError:
            seq = 1

    return f"{prefix}{seq:03d}"


def _get_available_process_types(template_name: str | None) -> List[str]:
    name = (template_name or "").strip()

    if "무지" in name:
        return ["CUT"]

    if "인쇄" in name:
        return ["CUT", "PRINT"]

    return ["CUT"]

    

def _get_primary_outsource_process_type(template_name: str | None) -> str:
    available_process_types = _get_available_process_types(template_name)

    if "PRINT" in available_process_types:
        return "PRINT"

    return "CUT"



def _is_bohyun_target_work_group(
    process_type: str,
    template_name: str | None,
) -> bool:
    available_process_types = _get_available_process_types(template_name)
    is_print_product = "PRINT" in available_process_types

    if process_type == "CUT":
        return not is_print_product

    if process_type == "PRINT":
        return is_print_product

    if process_type == "DIECUT":
        return True

    return False    

def _is_purchase_order_target_process(
    process_type: str,
    template_name: str | None,
) -> bool:
    available_process_types = _get_available_process_types(template_name)
    return process_type in available_process_types

def _get_inbound_partner_name(process_type: str, template_name: str) -> str:
    available = _get_available_process_types(template_name)

    is_print_product = "PRINT" in available

    if process_type == "CUT":
        return "상림" if is_print_product else "보현"

    if process_type == "PRINT":
        return "상림"

    return ""


def _build_instruction_out(
    db: Session,
    instruction: OutsourceWorkInstruction,
) -> OutsourceWorkInstructionOut:
    item_rows = (
        db.execute(
            select(OutsourceWorkInstructionItem, Lot, OrderLine, Product)
            .join(Lot, Lot.lot_id == OutsourceWorkInstructionItem.lot_id)
            .join(OrderLine, OrderLine.order_line_id == Lot.order_line_id)
            .join(Product, Product.product_id == Lot.product_id)
            .where(
                OutsourceWorkInstructionItem.outsource_work_instruction_id
                == instruction.outsource_work_instruction_id
            )
        )
        .all()
    )

    file_rows = (
        db.execute(
            select(OutsourceWorkInstructionFile).where(
                OutsourceWorkInstructionFile.outsource_work_instruction_id
                == instruction.outsource_work_instruction_id
            )
        )
        .scalars()
        .all()
    )

    return OutsourceWorkInstructionOut(
        outsource_work_instruction_id=instruction.outsource_work_instruction_id,
        instruction_no=instruction.instruction_no,
        instruction_date=instruction.instruction_date,
        process_type=instruction.process_type,
        partner_id=instruction.partner_id,
        is_bundle=instruction.is_bundle,
        memo=instruction.memo,
        created_at=instruction.created_at,
        updated_at=instruction.updated_at,
        items=[
            OutsourceWorkInstructionItemOut(
                outsource_work_instruction_item_id=item.outsource_work_instruction_item_id,
                lot_id=lot.lot_id,
                lot_no=lot.lot_no,
                order_no=order_line.order_no,
                line_no=order_line.line_no,
                product_code=product.product_code,
                product_name=product.product_name,
                lot_qty=lot.lot_qty,
                process_type=item.process_type,
            )
            for item, lot, order_line, product in item_rows
        ],
        files=[
            OutsourceWorkInstructionFileOut.model_validate(x, from_attributes=True)
            for x in file_rows
        ],
    )


def _create_instruction(
    db: Session,
    instruction_date: date,
    process_type: str,
    partner_id: int,
    lot_ids: list[int],
    memo: str | None,
    files: list,
    groups: list[OutsourceWorkInstructionGroupCreate] | None = None,
) -> OutsourceWorkInstruction:
    instruction = OutsourceWorkInstruction(
        instruction_no=_generate_instruction_no(db, instruction_date),
        instruction_date=instruction_date,
        process_type=process_type,
        partner_id=partner_id,
        is_bundle=len(lot_ids) > 1,
        memo=memo,
    )
    db.add(instruction)
    db.flush()

    if groups:
        _create_work_groups(
            db=db,
            instruction_id=instruction.outsource_work_instruction_id,
            instruction_date=instruction_date,
            process_type=process_type,
            groups=groups,
        )

    for lot_id in lot_ids:
        db.add(
            OutsourceWorkInstructionItem(
                outsource_work_instruction_id=instruction.outsource_work_instruction_id,
                lot_id=lot_id,
                process_type=process_type,
            )
        )

    for file in files:
        db.add(
            OutsourceWorkInstructionFile(
                outsource_work_instruction_id=instruction.outsource_work_instruction_id,
                file_name=file.file_name,
                file_path=file.file_path,
                content_type=file.content_type,
            )
        )

    db.flush()
    return instruction


def _filter_groups_for_lot_ids(
    groups: list[OutsourceWorkInstructionGroupCreate],
    allowed_lot_ids: list[int],
) -> list[OutsourceWorkInstructionGroupCreate]:
    allowed_set = set(allowed_lot_ids)
    filtered_groups: list[OutsourceWorkInstructionGroupCreate] = []

    for group in groups:
        filtered_items = [
            OutsourceWorkInstructionGroupItemCreate(
                lot_id=item.lot_id,
                cuts_per_sheet=item.cuts_per_sheet,
                expected_output_qty=item.expected_output_qty,
                remark=item.remark,
            )
            for item in group.items
            if item.lot_id in allowed_set
        ]

        if not filtered_items:
            continue

        sheet_cut_count = group.sheet_cut_count

        if group.is_bundle:
            sheet_cut_count = sum(item.cuts_per_sheet for item in filtered_items)

        filtered_groups.append(
            OutsourceWorkInstructionGroupCreate(
                group_seq=group.group_seq,
                is_bundle=group.is_bundle,
                sheet_qty=group.sheet_qty,
                length_m=group.length_m,
                sheet_cut_count=sheet_cut_count,
                remark=group.remark,
                items=filtered_items,
            )
        )

    return filtered_groups


def _utcnow() -> datetime:
    return datetime.now(timezone.utc)


def _generate_purchase_order_no(db: Session, process_type: str, order_date) -> str:
    normalized = (process_type or "").strip().upper()
    prefix = "OCUT" if normalized == "CUT" else "OPRT"
    ymd = order_date.strftime("%Y%m%d")
    like_prefix = f"{prefix}-{ymd}-"

    last_no = (
        db.execute(
            select(OutsourcePurchaseOrder.purchase_order_no)
            .where(OutsourcePurchaseOrder.purchase_order_no.like(f"{like_prefix}%"))
            .order_by(OutsourcePurchaseOrder.purchase_order_no.desc())
            .limit(1)
        )
        .scalar_one_or_none()
    )

    if last_no:
        try:
            seq = int(last_no.split("-")[-1]) + 1
        except ValueError:
            seq = 1
    else:
        seq = 1

    return f"{prefix}-{ymd}-{seq:03d}"

def _get_work_group_month_code(value: date) -> str:
    month_codes = {
        1: "A",
        2: "B",
        3: "C",
        4: "D",
        5: "E",
        6: "F",
        7: "G",
        8: "H",
        9: "I",
        10: "J",
        11: "K",
        12: "L",
    }

    return month_codes[value.month]


def _generate_work_group_seq(
    db: Session,
    instruction_date: date,
) -> str:
    prefix = f"{_get_work_group_month_code(instruction_date)}{instruction_date.day:02d}"

    existing_codes = (
        db.execute(
            select(OutsourceWorkGroup.group_seq)
            .join(
                OutsourceWorkInstruction,
                OutsourceWorkInstruction.outsource_work_instruction_id
                == OutsourceWorkGroup.outsource_work_instruction_id,
            )
            .where(OutsourceWorkInstruction.instruction_date == instruction_date)
            .where(OutsourceWorkGroup.group_seq.like(f"{prefix}%"))
            .order_by(OutsourceWorkGroup.group_seq.desc())
            .with_for_update()
        )
        .scalars()
        .all()
    )

    max_seq = 0

    for code in existing_codes:
        suffix = str(code).replace(prefix, "", 1)

        if not suffix.isdigit():
            continue

        max_seq = max(max_seq, int(suffix))

    return f"{prefix}{max_seq + 1}"


def _build_purchase_order_out(
    db: Session,
    purchase_order: OutsourcePurchaseOrder,
) -> OutsourcePurchaseOrderOut:
    item_rows = (
        db.execute(
            select(OutsourcePurchaseOrderItem, Lot, OrderLine, Product)
            .join(Lot, Lot.lot_id == OutsourcePurchaseOrderItem.lot_id)
            .join(OrderLine, OrderLine.order_line_id == Lot.order_line_id)
            .join(Product, Product.product_id == Lot.product_id)
            .where(
                OutsourcePurchaseOrderItem.outsource_purchase_order_id
                == purchase_order.outsource_purchase_order_id
            )
            .order_by(OutsourcePurchaseOrderItem.item_seq.asc())
        )
        .all()
    )

    items: list[OutsourcePurchaseOrderItemOut] = []

    for item, lot, order_line, product in item_rows:
        items.append(
            OutsourcePurchaseOrderItemOut(
                outsource_purchase_order_item_id=item.outsource_purchase_order_item_id,
                outsource_purchase_order_id=item.outsource_purchase_order_id,
                lot_id=item.lot_id,
                outsource_work_instruction_id=item.outsource_work_instruction_id,
                item_seq=item.item_seq,
                qty=item.qty,
                status=item.status,
                vendor_received_at=item.vendor_received_at,
                work_done_at=item.work_done_at,
                shipped_at=item.shipped_at,
                work_done_qty=item.work_done_qty,
                bad_qty=item.bad_qty,
                work_done_remark=item.work_done_remark,
                created_at=item.created_at,
                lot_no=lot.lot_no,
                order_no=order_line.order_no,
                line_no=order_line.line_no,
                product_code=product.product_code,
                product_name=product.product_name,
                lot_qty=lot.lot_qty,
            )
        )

    outsource_partner = db.get(Partner, purchase_order.outsource_partner_id)
    inbound_partner = (
        db.get(Partner, purchase_order.inbound_partner_id)
        if purchase_order.inbound_partner_id
        else None
    )

    return OutsourcePurchaseOrderOut(
        outsource_purchase_order_id=purchase_order.outsource_purchase_order_id,
        purchase_order_no=purchase_order.purchase_order_no,
        purchase_order_date=purchase_order.purchase_order_date,
        due_date=purchase_order.due_date,
        process_type=purchase_order.process_type,
        outsource_partner_id=purchase_order.outsource_partner_id,
        inbound_partner_id=purchase_order.inbound_partner_id,
        work_description=purchase_order.work_description,
        remark=purchase_order.remark,
        qty=purchase_order.qty,
        unit_price=purchase_order.unit_price,
        supply_amount=purchase_order.supply_amount,
        vat_amount=purchase_order.vat_amount,
        total_amount=purchase_order.total_amount,
        created_at=purchase_order.created_at,
        updated_at=purchase_order.updated_at,
        outsource_partner_name=outsource_partner.name if outsource_partner else None,
        inbound_partner_name=inbound_partner.name if inbound_partner else None,
        items=items,
    )


def _resolve_group_sheet_cut_count(
    db: Session,
    process_type: str,
    group_payload: OutsourceWorkInstructionGroupCreate,
) -> int:
    if group_payload.is_bundle:
        if group_payload.sheet_cut_count is None or group_payload.sheet_cut_count <= 0:
            raise HTTPException(
                status_code=409,
                detail="sheet_cut_count is required for bundle group",
            )

        cuts_sum = sum(item.cuts_per_sheet for item in group_payload.items)

        if cuts_sum != group_payload.sheet_cut_count:
            raise HTTPException(
                status_code=409,
                detail="sum(cuts_per_sheet) must equal sheet_cut_count for bundle group",
            )

        return group_payload.sheet_cut_count

    if len(group_payload.items) != 1:
        raise HTTPException(
            status_code=409,
            detail="non-bundle group must contain exactly one item",
        )

    lot_id = group_payload.items[0].lot_id

    cut_qty_per_panel = db.execute(
        select(Product.cut_qty_per_panel)
        .join(Lot, Lot.product_id == Product.product_id)
        .where(Lot.lot_id == lot_id)
    ).scalar_one_or_none()

    if cut_qty_per_panel is not None and int(cut_qty_per_panel) > 0:
        return int(cut_qty_per_panel)

    if group_payload.sheet_cut_count is not None and group_payload.sheet_cut_count > 0:
        return group_payload.sheet_cut_count

    raise HTTPException(
        status_code=409,
        detail="cut_qty_per_panel or sheet_cut_count is required for non-bundle group",
    )


def _create_work_groups(
    db: Session,
    instruction_id: int,
    instruction_date: date,
    process_type: str,
    groups: list[OutsourceWorkInstructionGroupCreate],
) -> None:
    for group_payload in groups:
        sheet_cut_count = _resolve_group_sheet_cut_count(
            db=db,
            process_type=process_type,
            group_payload=group_payload,
        )

        work_group = OutsourceWorkGroup(
            outsource_work_instruction_id=instruction_id,
            group_seq=_generate_work_group_seq(db, instruction_date),
            process_type=process_type,
            is_bundle=group_payload.is_bundle,
            sheet_qty=group_payload.sheet_qty,
            length_m=group_payload.length_m,
            sheet_cut_count=sheet_cut_count,
            remark=group_payload.remark,
        )


@router.get("/bohyun-groups", response_model=BohyunOutsourceGroupListOut)
def get_bohyun_outsource_groups(
    date_from: date | None = Query(default=None),
    date_to: date | None = Query(default=None),
    process_type: str | None = Query(default=None),
    status: str | None = Query(default=None),
    q: str | None = Query(default=None),
    db: Session = Depends(get_db),
):
    if process_type and process_type not in ("CUT", "PRINT", "DIECUT"):
        raise HTTPException(status_code=409, detail="Invalid process_type")

    if status and status not in (
        BOHYUN_UI_STATUS_WAITING_INBOUND,
        BOHYUN_UI_STATUS_INBOUNDED,
        BOHYUN_UI_STATUS_WORK_DONE,
        BOHYUN_UI_STATUS_SHIPPED,
    ):
        raise HTTPException(status_code=409, detail="Invalid status")

    stmt = (
        select(OutsourceWorkGroup, OutsourceWorkInstruction, Partner)
        .join(
            OutsourceWorkInstruction,
            OutsourceWorkInstruction.outsource_work_instruction_id
            == OutsourceWorkGroup.outsource_work_instruction_id,
        )
        .join(Partner, Partner.partner_id == OutsourceWorkInstruction.partner_id)
        .where(OutsourceWorkGroup.process_type.in_(("CUT", "PRINT", "DIECUT")))
        .order_by(
            OutsourceWorkInstruction.instruction_date.desc(),
            OutsourceWorkInstruction.instruction_no.desc(),
            OutsourceWorkGroup.group_seq.asc(),
        )
    )

    if date_from:
        stmt = stmt.where(OutsourceWorkInstruction.instruction_date >= date_from)

    if date_to:
        stmt = stmt.where(OutsourceWorkInstruction.instruction_date <= date_to)

    if process_type:
        stmt = stmt.where(OutsourceWorkGroup.process_type == process_type)

    if status == BOHYUN_UI_STATUS_WAITING_INBOUND:
        stmt = stmt.where(OutsourceWorkGroup.status.is_(None))
    elif status == BOHYUN_UI_STATUS_INBOUNDED:
        stmt = stmt.where(OutsourceWorkGroup.status == BOHYUN_DB_STATUS_VENDOR_RECEIVED)
    elif status == BOHYUN_UI_STATUS_WORK_DONE:
        stmt = stmt.where(OutsourceWorkGroup.status == BOHYUN_DB_STATUS_WORK_DONE)
    elif status == BOHYUN_UI_STATUS_SHIPPED:
        stmt = stmt.where(OutsourceWorkGroup.status == BOHYUN_DB_STATUS_SHIPPED)
    else:
        stmt = stmt.where(
            or_(
                OutsourceWorkGroup.status.is_(None),
                OutsourceWorkGroup.status != BOHYUN_DB_STATUS_SHIPPED,
            )
        )

    if q:
        like = f"%{q.strip()}%"

        exists_item_stmt = (
            select(OutsourceWorkGroupItem.outsource_work_group_item_id)
            .join(Lot, Lot.lot_id == OutsourceWorkGroupItem.lot_id)
            .join(OrderLine, OrderLine.order_line_id == Lot.order_line_id)
            .join(Product, Product.product_id == Lot.product_id)
            .where(
                OutsourceWorkGroupItem.outsource_work_group_id
                == OutsourceWorkGroup.outsource_work_group_id
            )
            .where(
                (Lot.lot_no.like(like))
                | (OrderLine.order_no.like(like))
                | (Product.product_code.like(like))
                | (Product.product_name.like(like))
            )
            .limit(1)
        )

        stmt = stmt.where(
            (OutsourceWorkInstruction.instruction_no.like(like))
            | (Partner.name.like(like))
            | exists(exists_item_stmt)
        )

    rows = db.execute(stmt).all()

    result_items: list[BohyunOutsourceGroupListItemOut] = []

    for work_group, instruction, partner in rows:
        group_item_rows = (
            db.execute(
                select(
                    OutsourceWorkGroupItem,
                    Lot,
                    OrderLine,
                    Product,
                    RoutingTemplate,
                )
                .join(Lot, Lot.lot_id == OutsourceWorkGroupItem.lot_id)
                .join(OrderLine, OrderLine.order_line_id == Lot.order_line_id)
                .join(Product, Product.product_id == Lot.product_id)
                .join(
                    RoutingTemplate,
                    RoutingTemplate.routing_template_id == Product.routing_template_id,
                )
                .where(
                    OutsourceWorkGroupItem.outsource_work_group_id
                    == work_group.outsource_work_group_id
                )
                .order_by(
                    Lot.lot_no.asc(),
                    OutsourceWorkGroupItem.outsource_work_group_item_id.asc(),
                )
            )
            .all()
        )

        if not group_item_rows:
            continue

        target_group_item_rows = [
            row
            for row in group_item_rows
            if _is_bohyun_target_work_group(
                work_group.process_type,
                row[4].template_name,
            )
        ]

        if not target_group_item_rows:
            continue

        group_items: list[BohyunOutsourceGroupItemOut] = []
        lot_nos: list[str] = []
        product_names: list[str] = []

        for group_item, lot, order_line, product, routing_template in target_group_item_rows:
            if lot.lot_no:
                lot_nos.append(lot.lot_no)

            if product.product_name:
                product_names.append(product.product_name)

            group_items.append(
                BohyunOutsourceGroupItemOut(
                    outsource_work_group_item_id=group_item.outsource_work_group_item_id,
                    lot_id=lot.lot_id,
                    lot_no=lot.lot_no,
                    order_no=order_line.order_no,
                    line_no=order_line.line_no,
                    product_id=product.product_id,
                    product_code=product.product_code,
                    product_name=product.product_name,
                    cuts_per_sheet=group_item.cuts_per_sheet,
                    expected_output_qty=group_item.expected_output_qty,
                    actual_output_qty=group_item.actual_output_qty,
                    remark=group_item.remark,
                )
            )

        result_items.append(
            BohyunOutsourceGroupListItemOut(
                outsource_work_group_id=work_group.outsource_work_group_id,
                outsource_work_instruction_id=instruction.outsource_work_instruction_id,
                instruction_no=instruction.instruction_no,
                instruction_date=instruction.instruction_date,
                process_type=work_group.process_type,
                partner_id=partner.partner_id,
                partner_name=partner.name,
                inbound_source_name=_get_bohyun_inbound_source_name(work_group.process_type),
                is_bundle=work_group.is_bundle,
                group_seq=work_group.group_seq,
                sheet_qty=work_group.sheet_qty,
                work_done_sheet_qty=work_group.work_done_sheet_qty,
                length_m=work_group.length_m,
                sheet_cut_count=work_group.sheet_cut_count,
                status=_to_bohyun_ui_status(work_group.status),
                vendor_received_at=work_group.vendor_received_at,
                work_done_at=work_group.work_done_at,
                shipped_at=work_group.shipped_at,
                outsource_processing_fee=work_group.outsource_processing_fee,
                work_done_remark=work_group.work_done_remark,
                lot_nos=lot_nos,
                product_names=list(dict.fromkeys(product_names)),
                items=group_items,
            )
        )

    return BohyunOutsourceGroupListOut(
        items=result_items,
        total_count=len(result_items),
    )


@router.post("/bohyun-groups/ship-batch")
def ship_bohyun_outsource_groups(
    payload: BohyunOutsourceGroupShipBatch,
    db: Session = Depends(get_db),
):
    requested_group_ids = list(dict.fromkeys(payload.group_ids))

    work_groups = (
        db.execute(
            select(OutsourceWorkGroup)
            .where(OutsourceWorkGroup.outsource_work_group_id.in_(requested_group_ids))
        )
        .scalars()
        .all()
    )

    if len(work_groups) != len(requested_group_ids):
        raise HTTPException(
            status_code=404,
            detail="Some outsource work groups were not found",
        )

    invalid_groups = [
        work_group
        for work_group in work_groups
        if work_group.status != BOHYUN_DB_STATUS_WORK_DONE
    ]

    if invalid_groups:
        raise HTTPException(
            status_code=409,
            detail="Only work done groups can be shipped",
        )

    now = datetime.now()

    for work_group in work_groups:
        work_group.status = BOHYUN_DB_STATUS_SHIPPED
        work_group.shipped_at = now

    db.commit()

    return {"success": True}


@router.post("/bohyun-groups/{group_id}/inbound")
def inbound_bohyun_outsource_group(
    group_id: int,
    db: Session = Depends(get_db),
):
    work_group = db.get(OutsourceWorkGroup, group_id)

    if not work_group:
        raise HTTPException(
            status_code=404,
            detail="Outsource work group not found",
        )

    if work_group.status == BOHYUN_DB_STATUS_SHIPPED:
        raise HTTPException(
            status_code=409,
            detail="Already shipped group cannot be inbounded",
        )

    if work_group.status == BOHYUN_DB_STATUS_WORK_DONE:
        raise HTTPException(
            status_code=409,
            detail="Already work done group cannot be inbounded",
        )

    if work_group.status == BOHYUN_DB_STATUS_VENDOR_RECEIVED:
        raise HTTPException(
            status_code=409,
            detail="Already inbounded",
        )

    work_group.status = BOHYUN_DB_STATUS_VENDOR_RECEIVED
    work_group.vendor_received_at = datetime.now()

    db.commit()

    return {"success": True}


@router.post("/bohyun-groups/{group_id}/work-done")
def complete_bohyun_outsource_group_work(
    group_id: int,
    payload: BohyunOutsourceGroupWorkDone,
    db: Session = Depends(get_db),
):
    work_group = db.get(OutsourceWorkGroup, group_id)

    if not work_group:
        raise HTTPException(
            status_code=404,
            detail="Outsource work group not found",
        )

    if work_group.status == BOHYUN_DB_STATUS_SHIPPED:
        raise HTTPException(
            status_code=409,
            detail="Already shipped group cannot be work done",
        )

    if work_group.status == BOHYUN_DB_STATUS_WORK_DONE:
        raise HTTPException(
            status_code=409,
            detail="Already work done",
        )

    if work_group.status != BOHYUN_DB_STATUS_VENDOR_RECEIVED:
        raise HTTPException(
            status_code=409,
            detail="Only inbounded group can be work done",
        )

    group_items = (
        db.execute(
            select(OutsourceWorkGroupItem)
            .where(
                OutsourceWorkGroupItem.outsource_work_group_id
                == work_group.outsource_work_group_id
            )
            .order_by(OutsourceWorkGroupItem.outsource_work_group_item_id.asc())
        )
        .scalars()
        .all()
    )

    if not group_items:
        raise HTTPException(
            status_code=409,
            detail="Outsource work group has no items",
        )

    for group_item in group_items:
        group_item.actual_output_qty = payload.work_done_sheet_qty * group_item.cuts_per_sheet

    work_group.status = BOHYUN_DB_STATUS_WORK_DONE
    work_group.work_done_at = datetime.now()
    work_group.work_done_sheet_qty = payload.work_done_sheet_qty
    work_group.outsource_processing_fee = payload.outsource_processing_fee
    work_group.work_done_remark = payload.remark

    db.commit()

    return {"success": True}


@router.get("/purchase-orders/{outsource_purchase_order_id}/excel")
def download_outsource_purchase_order_excel(
    outsource_purchase_order_id: int,
    db: Session = Depends(get_db),
):
    purchase_order = db.get(OutsourcePurchaseOrder, outsource_purchase_order_id)

    if not purchase_order:
        raise HTTPException(status_code=404, detail="Outsource purchase order not found")

    result = _build_purchase_order_out(db, purchase_order)

    if purchase_order.process_type == "PRINT":
        file_bytes = _build_print_purchase_order_excel_template_bytes(
            result,
            purchase_order.form_snapshot_json,
        )
    else:
        file_bytes = _build_purchase_order_excel_template_bytes(
            result,
            purchase_order.form_snapshot_json,
        )

    filename = f"{result.purchase_order_no}.xlsx"

    return StreamingResponse(
        BytesIO(file_bytes),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
        },
    )


@lru_cache(maxsize=1)
def _get_cut_template_bytes() -> bytes:
    template_path = (
        Path(__file__).resolve().parents[2]
        / "templates"
        / "outsource_purchase_order_cut_template.xlsx"
    )
    return template_path.read_bytes()


@lru_cache(maxsize=1)
def _get_print_template_bytes() -> bytes:
    template_path = (
        Path(__file__).resolve().parents[2]
        / "templates"
        / "outsource_purchase_order_print_template.xlsx"
    )
    return template_path.read_bytes()


def _build_print_purchase_order_excel_template_bytes(
    purchase_order: OutsourcePurchaseOrderOut,
    form_snapshot: dict | None,
) -> bytes:
    wb = load_workbook(BytesIO(_get_print_template_bytes()))
    ws = wb.active
    merged_map = _build_merged_cell_map(ws)

    snapshot = form_snapshot or {}
    rows = snapshot.get("rows") or []

    _set_merged_safe(ws, merged_map, "D5", snapshot.get("request_company_name") or "")
    _set_merged_safe(ws, merged_map, "F5", snapshot.get("requester_name") or "")
    _set_merged_safe(
        ws,
        merged_map,
        "J5",
        snapshot.get("purchase_order_date") or str(purchase_order.purchase_order_date),
    )

    start_row = 8
    max_rows = 16

    for idx, row_data in enumerate(rows[:max_rows]):
        r = start_row + idx

        _set_merged_safe(ws, merged_map, f"B{r}", row_data.get("no", ""))
        _set_merged_safe(ws, merged_map, f"C{r}", row_data.get("customer_name", ""))
        _set_merged_safe(ws, merged_map, f"D{r}", row_data.get("product_name", ""))
        _set_merged_safe(ws, merged_map, f"E{r}", row_data.get("material_spec", ""))
        _set_merged_safe(ws, merged_map, f"F{r}", row_data.get("print_sheet_qty", ""))
        _set_merged_safe(ws, merged_map, f"G{r}", row_data.get("sample", ""))
        _set_merged_safe(ws, merged_map, f"H{r}", row_data.get("plate_count", ""))
        _set_merged_safe(ws, merged_map, f"I{r}", row_data.get("color_name", ""))
        _set_merged_safe(ws, merged_map, f"J{r}", row_data.get("material_type", ""))
        _set_merged_safe(ws, merged_map, f"K{r}", row_data.get("remark", ""))

    _set_merged_safe(ws, merged_map, "D24", snapshot.get("footer_remark") or "")

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream.getvalue()


def _build_merged_cell_map(ws) -> dict[str, str]:
    merged_map: dict[str, str] = {}

    for merged_range in ws.merged_cells.ranges:
        start_ref = merged_range.start_cell.coordinate

        for row in ws[merged_range.coord]:
            for cell in row:
                merged_map[cell.coordinate] = start_ref

    return merged_map


def _set_merged_safe(ws, merged_map: dict[str, str], cell_ref: str, value) -> None:
    target_ref = merged_map.get(cell_ref, cell_ref)
    ws[target_ref] = value


def _build_purchase_order_excel_template_bytes(
    purchase_order: OutsourcePurchaseOrderOut,
    form_snapshot: dict | None,
) -> bytes:
    wb = load_workbook(BytesIO(_get_cut_template_bytes()))
    ws = wb.active
    merged_map = _build_merged_cell_map(ws)

    snapshot = form_snapshot or {}
    rows = snapshot.get("rows") or []

    _set_merged_safe(ws, merged_map, "E5", snapshot.get("request_company_name") or "")
    _set_merged_safe(ws, merged_map, "I5", snapshot.get("requester_name") or "")
    _set_merged_safe(
        ws,
        merged_map,
        "P5",
        snapshot.get("purchase_order_date") or str(purchase_order.purchase_order_date),
    )
    _set_merged_safe(
        ws,
        merged_map,
        "H7",
        snapshot.get("raw_material_inbound_text") or "",
    )

    start_row = 10
    max_rows = 16

    for idx, row_data in enumerate(rows[:max_rows]):
        r = start_row + idx

        _set_merged_safe(ws, merged_map, f"B{r}", row_data.get("no", ""))
        _set_merged_safe(ws, merged_map, f"C{r}", row_data.get("raw_material_text", ""))
        _set_merged_safe(ws, merged_map, f"H{r}", row_data.get("length_m_text", ""))
        _set_merged_safe(ws, merged_map, f"I{r}", row_data.get("inbound_place_text", ""))
        _set_merged_safe(ws, merged_map, f"O{r}", row_data.get("cut_spec_text", ""))
        _set_merged_safe(ws, merged_map, f"P{r}", row_data.get("sheet_qty_text", ""))

    _set_merged_safe(ws, merged_map, "H30", snapshot.get("stock_500_width_text") or "")
    _set_merged_safe(ws, merged_map, "J30", snapshot.get("stock_600_width_text") or "")
    _set_merged_safe(ws, merged_map, "O30", snapshot.get("stock_600_tpt0268_text") or "")
    _set_merged_safe(ws, merged_map, "B32", snapshot.get("remark") or "")

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream.getvalue()


@router.post(
    "/upload-plate-data",
    response_model=OutsourceWorkInstructionPlateUploadOut,
    status_code=http_status.HTTP_201_CREATED,
)
async def upload_plate_data(
    file: UploadFile = File(...),
):
    uploaded_at = datetime.now()
    content = await file.read()
    size_bytes = len(content)

    _validate_plate_data_upload(file, size_bytes)

    target_path = _build_plate_data_path(file.filename or "plate_data.bin", uploaded_at)

    try:
        target_path.write_bytes(content)
    except OSError:
        raise HTTPException(status_code=500, detail="Failed to save uploaded file")

    return OutsourceWorkInstructionPlateUploadOut(
        file_name=Path(file.filename or target_path.name).name,
        file_path=str(target_path),
        content_type=file.content_type,
        file_size=size_bytes,
        uploaded_at=uploaded_at,
    )


@router.get("/candidates", response_model=OutsourceWorkInstructionCandidateLotListOut)
def get_candidate_lots(
    process_type: str | None = Query(default=None),
    q: str | None = Query(default=None),
    db: Session = Depends(get_db),
):
    if process_type and process_type not in ("CUT", "PRINT"):
        raise HTTPException(status_code=409, detail="Invalid process_type")

    stmt = (
        select(Lot, OrderLine, Product, Partner, RoutingTemplate)
        .join(OrderLine, OrderLine.order_line_id == Lot.order_line_id)
        .join(Product, Product.product_id == Lot.product_id)
        .join(Partner, Partner.partner_id == OrderLine.partner_id)
        .join(RoutingTemplate, RoutingTemplate.routing_template_id == Product.routing_template_id)
        .order_by(Lot.created_date.desc(), Lot.lot_no.asc())
    )

    if q:
        like = f"%{q.strip()}%"
        stmt = stmt.where(
            (Lot.lot_no.like(like))
            | (OrderLine.order_no.like(like))
            | (Product.product_code.like(like))
            | (Product.product_name.like(like))
            | (Partner.name.like(like))
        )

    rows = db.execute(stmt).all()
    items: list[OutsourceWorkInstructionCandidateLotOut] = []

    for lot, order_line, product, partner, routing_template in rows:
        available = _get_available_process_types(routing_template.template_name)

        if process_type and process_type not in available:
            continue

        required_processes = [process_type] if process_type else available

        exists_registered = db.execute(
            select(OutsourceWorkInstructionItem.outsource_work_instruction_item_id)
            .where(
                OutsourceWorkInstructionItem.lot_id == lot.lot_id,
                OutsourceWorkInstructionItem.process_type.in_(required_processes),
            )
            .limit(1)
        ).scalar_one_or_none()

        if exists_registered:
            continue

        items.append(
            OutsourceWorkInstructionCandidateLotOut(
                lot_id=lot.lot_id,
                lot_no=lot.lot_no,
                order_line_id=order_line.order_line_id,
                order_no=order_line.order_no,
                line_no=order_line.line_no,
                product_id=product.product_id,
                product_code=product.product_code,
                product_name=product.product_name,
                customer_partner_id=partner.partner_id,
                customer_partner_name=partner.name,
                lot_qty=lot.lot_qty,
                available_process_types=available,
                panel_width_mm=product.panel_width_mm,
                panel_length_mm=product.panel_length_mm,
                product_spec=product.product_spec,
                cut_qty_per_panel=product.cut_qty_per_panel,
            )
        )

    return OutsourceWorkInstructionCandidateLotListOut(items=items)


@router.post("", response_model=OutsourceWorkInstructionOut, status_code=http_status.HTTP_201_CREATED)
def create_outsource_work_instruction(
    payload: OutsourceWorkInstructionCreate,
    db: Session = Depends(get_db),
):
    if payload.process_type not in ("CUT", "PRINT"):
        raise HTTPException(status_code=409, detail="Invalid process_type")

    partner = db.get(Partner, payload.partner_id)

    if not partner or not partner.is_active:
        raise HTTPException(status_code=404, detail="Partner not found or inactive")

    if len(payload.lot_ids) > 1 and len(payload.files) > 1:
        raise HTTPException(
            status_code=409,
            detail="Bundle work instruction allows only one plate data file",
        )

    if payload.process_type == "PRINT" and len(payload.files) == 0:
        raise HTTPException(
            status_code=409,
            detail="PRINT work instruction requires plate data file",
        )

    lots = (
        db.execute(
            select(Lot, OrderLine, Product, RoutingTemplate)
            .join(OrderLine, OrderLine.order_line_id == Lot.order_line_id)
            .join(Product, Product.product_id == Lot.product_id)
            .join(RoutingTemplate, RoutingTemplate.routing_template_id == Product.routing_template_id)
            .where(Lot.lot_id.in_(payload.lot_ids))
        )
        .all()
    )

    if len(lots) != len(set(payload.lot_ids)):
        raise HTTPException(status_code=404, detail="Some lots were not found")

    for lot, order_line, product, routing_template in lots:
        available = _get_available_process_types(routing_template.template_name)

        if payload.process_type not in available:
            raise HTTPException(
                status_code=409,
                detail=f"LOT {lot.lot_no} cannot be registered for process {payload.process_type}",
            )

        exists_registered = db.execute(
            select(OutsourceWorkInstructionItem.outsource_work_instruction_item_id)
            .where(
                OutsourceWorkInstructionItem.lot_id == lot.lot_id,
                OutsourceWorkInstructionItem.process_type == payload.process_type,
            )
            .limit(1)
        ).scalar_one_or_none()

        if exists_registered:
            raise HTTPException(
                status_code=409,
                detail=f"LOT {lot.lot_no} is already registered for process {payload.process_type}",
            )

    instruction = OutsourceWorkInstruction(
        instruction_no=_generate_instruction_no(db, payload.instruction_date),
        instruction_date=payload.instruction_date,
        process_type=payload.process_type,
        partner_id=payload.customer_partner_id,
        is_bundle=len(payload.lot_ids) > 1,
        memo=payload.memo,
    )
    db.add(instruction)
    db.flush()

    if payload.groups:
        _create_work_groups(
            db=db,
            instruction_id=instruction.outsource_work_instruction_id,
            instruction_date=payload.instruction_date,
            process_type=payload.process_type,
            groups=payload.groups,
        )

    for lot_id in payload.lot_ids:
        db.add(
            OutsourceWorkInstructionItem(
                outsource_work_instruction_id=instruction.outsource_work_instruction_id,
                lot_id=lot_id,
                process_type=payload.process_type,
            )
        )

    for file in payload.files:
        db.add(
            OutsourceWorkInstructionFile(
                outsource_work_instruction_id=instruction.outsource_work_instruction_id,
                file_name=file.file_name,
                file_path=file.file_path,
                content_type=file.content_type,
            )
        )

    db.commit()
    db.refresh(instruction)

    return _build_instruction_out(db, instruction)


@router.post(
    "/batch",
    response_model=OutsourceWorkInstructionBatchOut,
    status_code=http_status.HTTP_201_CREATED,
)
def create_outsource_work_instruction_batch(
    payload: OutsourceWorkInstructionBatchCreate,
    db: Session = Depends(get_db),
):
    created_instructions: list[OutsourceWorkInstruction] = []

    for group in payload.groups:
        partner = db.get(Partner, group.customer_partner_id)

        if not partner or not partner.is_active:
            raise HTTPException(status_code=404, detail="Partner not found or inactive")

        if len(group.lot_ids) > 1 and len(group.files) > 1:
            raise HTTPException(
                status_code=409,
                detail="Bundle work instruction allows only one plate data file",
            )

        if len(group.lot_ids) > 1 and len(group.files) == 0:
            raise HTTPException(
                status_code=409,
                detail="Bundle work instruction requires plate data file",
            )

        lots = (
            db.execute(
                select(Lot, OrderLine, Product, RoutingTemplate)
                .join(OrderLine, OrderLine.order_line_id == Lot.order_line_id)
                .join(Product, Product.product_id == Lot.product_id)
                .join(RoutingTemplate, RoutingTemplate.routing_template_id == Product.routing_template_id)
                .where(Lot.lot_id.in_(group.lot_ids))
            )
            .all()
        )

        if len(lots) != len(set(group.lot_ids)):
            raise HTTPException(status_code=404, detail="Some lots were not found")

        cut_lot_ids: list[int] = []
        print_lot_ids: list[int] = []

        for lot, order_line, product, routing_template in lots:
            primary_process_type = _get_primary_outsource_process_type(
                routing_template.template_name
            )

            if primary_process_type == "PRINT":
                print_lot_ids.append(lot.lot_id)
            else:
                cut_lot_ids.append(lot.lot_id)

        if not cut_lot_ids and not print_lot_ids:
            raise HTTPException(status_code=409, detail="No available process for selected lots")

        if cut_lot_ids:
            for lot_id in cut_lot_ids:
                exists_registered = db.execute(
                    select(OutsourceWorkInstructionItem.outsource_work_instruction_item_id)
                    .where(
                        OutsourceWorkInstructionItem.lot_id == lot_id,
                        OutsourceWorkInstructionItem.process_type == "CUT",
                    )
                    .limit(1)
                ).scalar_one_or_none()

                if exists_registered:
                    raise HTTPException(
                        status_code=409,
                        detail="Some lots are already registered for process CUT",
                    )

        if print_lot_ids:
            for lot_id in print_lot_ids:
                exists_registered = db.execute(
                    select(OutsourceWorkInstructionItem.outsource_work_instruction_item_id)
                    .where(
                        OutsourceWorkInstructionItem.lot_id == lot_id,
                        OutsourceWorkInstructionItem.process_type == "PRINT",
                    )
                    .limit(1)
                ).scalar_one_or_none()

                if exists_registered:
                    raise HTTPException(
                        status_code=409,
                        detail="Some lots are already registered for process PRINT",
                    )

        cut_groups = _filter_groups_for_lot_ids(group.groups, cut_lot_ids) if group.groups else []
        print_groups = _filter_groups_for_lot_ids(group.groups, print_lot_ids) if group.groups else []

        if cut_lot_ids:
            created_instructions.append(
                _create_instruction(
                    db=db,
                    instruction_date=payload.instruction_date,
                    process_type="CUT",
                    partner_id=group.customer_partner_id,
                    lot_ids=cut_lot_ids,
                    memo=group.memo,
                    files=[],
                    groups=cut_groups,
                )
            )

        if print_lot_ids:
            created_instructions.append(
                _create_instruction(
                    db=db,
                    instruction_date=payload.instruction_date,
                    process_type="PRINT",
                    partner_id=group.customer_partner_id,
                    lot_ids=print_lot_ids,
                    memo=group.memo,
                    files=group.files,
                    groups=print_groups,
                )
            )

    db.commit()

    for instruction in created_instructions:
        db.refresh(instruction)

    return OutsourceWorkInstructionBatchOut(
        items=[
            _build_instruction_out(db, instruction)
            for instruction in created_instructions
        ]
    )


@router.get(
    "/purchase-order-targets",
    response_model=OutsourcePurchaseOrderTargetListOut,
)
def get_purchase_order_targets(
    process_type: str,
    db: Session = Depends(get_db),
):
    normalized_process_type = (process_type or "").strip().upper()

    if normalized_process_type not in {"CUT", "PRINT"}:
        raise HTTPException(
            status_code=400,
            detail="process_type must be CUT or PRINT",
        )

    outsource_partner = _require_outsource_partner_by_process_type(
        db=db,
        process_type=normalized_process_type,
    )

    order_partner = aliased(Partner)

    rows = (
        db.execute(
            select(
                OutsourceWorkInstruction,
                OutsourceWorkInstructionItem,
                OutsourceWorkGroup,
                OutsourceWorkGroupItem,
                Lot,
                OrderLine,
                Product,
                Partner,
                RoutingTemplate,
                order_partner,
            )
            .join(
                OutsourceWorkInstructionItem,
                OutsourceWorkInstructionItem.outsource_work_instruction_id
                == OutsourceWorkInstruction.outsource_work_instruction_id,
            )
            .join(
                OutsourceWorkGroup,
                OutsourceWorkGroup.outsource_work_instruction_id
                == OutsourceWorkInstruction.outsource_work_instruction_id,
            )
            .join(
                OutsourceWorkGroupItem,
                OutsourceWorkGroupItem.outsource_work_group_id
                == OutsourceWorkGroup.outsource_work_group_id,
            )
            .join(Lot, Lot.lot_id == OutsourceWorkInstructionItem.lot_id)
            .join(OrderLine, OrderLine.order_line_id == Lot.order_line_id)
            .join(Product, Product.product_id == Lot.product_id)
            .join(RoutingTemplate, RoutingTemplate.routing_template_id == Product.routing_template_id)
            .join(order_partner, order_partner.partner_id == OrderLine.partner_id)
            .join(Partner, Partner.partner_id == OutsourceWorkInstruction.partner_id)
            .where(
                OutsourceWorkGroupItem.lot_id == Lot.lot_id,
                ~exists(
                    select(1)
                    .select_from(OutsourcePurchaseOrderItem)
                    .join(
                        OutsourcePurchaseOrder,
                        OutsourcePurchaseOrder.outsource_purchase_order_id
                        == OutsourcePurchaseOrderItem.outsource_purchase_order_id,
                    )
                    .where(
                        OutsourcePurchaseOrderItem.lot_id == Lot.lot_id,
                        OutsourcePurchaseOrder.process_type == normalized_process_type,
                    )
                ),
            )
            .order_by(
                OutsourceWorkInstruction.instruction_date.desc(),
                OutsourceWorkInstruction.instruction_no.desc(),
                Lot.lot_no.asc(),
            )
        )
        .all()
    )

    instruction_ids = list(
        {
            instruction.outsource_work_instruction_id
            for instruction, _, _, _, _, _, _, _, _, _ in rows
        }
    )

    file_map: dict[int, list[OutsourceWorkInstructionFileOut]] = {}

    if instruction_ids:
        file_rows = (
            db.execute(
                select(OutsourceWorkInstructionFile).where(
                    OutsourceWorkInstructionFile.outsource_work_instruction_id.in_(instruction_ids)
                )
            )
            .scalars()
            .all()
        )

        for file_row in file_rows:
            instruction_id = file_row.outsource_work_instruction_id

            if instruction_id not in file_map:
                file_map[instruction_id] = []

            file_map[instruction_id].append(
                OutsourceWorkInstructionFileOut.model_validate(
                    file_row,
                    from_attributes=True,
                )
            )

    items: list[OutsourcePurchaseOrderTargetOut] = []

    for (
        instruction,
        item,
        work_group,
        work_group_item,
        lot,
        order_line,
        product,
        instruction_partner,
        routing_template,
        source_partner,
    ) in rows:
        if not _is_purchase_order_target_process(
            normalized_process_type,
            routing_template.template_name,
        ):
            continue

        inbound_partner_name = _get_inbound_partner_name(
            normalized_process_type,
            routing_template.template_name,
        )
        available = _get_available_process_types(routing_template.template_name)
        is_print_product = "PRINT" in available

        items.append(
            OutsourcePurchaseOrderTargetOut(
                outsource_work_instruction_id=instruction.outsource_work_instruction_id,
                outsource_work_instruction_item_id=item.outsource_work_instruction_item_id,
                instruction_no=instruction.instruction_no,
                instruction_date=instruction.instruction_date,
                process_type=normalized_process_type,
                lot_id=lot.lot_id,
                lot_no=lot.lot_no,
                is_rework=lot.parent_lot_id is not None,
                order_line_id=order_line.order_line_id,
                order_no=order_line.order_no,
                line_no=order_line.line_no,
                product_id=product.product_id,
                product_code=product.product_code,
                product_name=product.product_name,
                customer_partner_id=source_partner.partner_id,
                customer_partner_name=source_partner.name,
                lot_qty=lot.lot_qty,
                outsource_partner_id=outsource_partner.partner_id,
                outsource_partner_name=outsource_partner.name,
                inbound_partner_name=inbound_partner_name,
                is_bundle=instruction.is_bundle,
                memo=instruction.memo,
                files=file_map.get(instruction.outsource_work_instruction_id, []),
                panel_width_mm=product.panel_width_mm,
                panel_length_mm=product.panel_length_mm,
                product_spec=product.product_spec,
                cut_qty_per_panel=product.cut_qty_per_panel,
                length_m=work_group.length_m,
                sheet_qty=work_group.sheet_qty,
                is_print_product=is_print_product,
            )
        )

    return OutsourcePurchaseOrderTargetListOut(items=items)


@router.post(
    "/purchase-orders",
    response_model=OutsourcePurchaseOrderOut,
    status_code=http_status.HTTP_201_CREATED,
)
def create_outsource_purchase_order(
    payload: OutsourcePurchaseOrderCreate,
    db: Session = Depends(get_db),
):
    normalized_process_type = (payload.process_type or "").strip().upper()

    if normalized_process_type not in ("CUT", "PRINT"):
        raise HTTPException(status_code=409, detail="Invalid process_type")

    outsource_partner = _require_outsource_partner_by_process_type(
        db=db,
        process_type=normalized_process_type,
    )

    if payload.inbound_partner_id:
        inbound_partner = db.get(Partner, payload.inbound_partner_id)

        if not inbound_partner or not inbound_partner.is_active:
            raise HTTPException(
                status_code=404,
                detail="Inbound partner not found or inactive",
            )

    lot_ids = [item.lot_id for item in payload.items]

    if len(lot_ids) != len(set(lot_ids)):
        raise HTTPException(
            status_code=409,
            detail="Duplicate lot exists in purchase order items",
        )

    lot_count = (
        db.execute(
            select(func.count())
            .select_from(Lot)
            .where(Lot.lot_id.in_(lot_ids))
        )
        .scalar_one()
    )

    if int(lot_count) != len(set(lot_ids)):
        raise HTTPException(status_code=409, detail="Some lots do not exist")

    already_ordered_lot_id = (
        db.execute(
            select(OutsourcePurchaseOrderItem.lot_id)
            .join(
                OutsourcePurchaseOrder,
                OutsourcePurchaseOrder.outsource_purchase_order_id
                == OutsourcePurchaseOrderItem.outsource_purchase_order_id,
            )
            .where(OutsourcePurchaseOrder.process_type == normalized_process_type)
            .where(OutsourcePurchaseOrderItem.lot_id.in_(lot_ids))
            .limit(1)
        )
        .scalar_one_or_none()
    )

    if already_ordered_lot_id is not None:
        raise HTTPException(
            status_code=409,
            detail=(
                f"LOT is already purchase ordered for process "
                f"{normalized_process_type}: lot_id={already_ordered_lot_id}"
            ),
        )

    lot_rows = (
        db.execute(
            select(Lot, Product, RoutingTemplate)
            .join(Product, Product.product_id == Lot.product_id)
            .join(
                RoutingTemplate,
                RoutingTemplate.routing_template_id == Product.routing_template_id,
            )
            .where(Lot.lot_id.in_(lot_ids))
        )
        .all()
    )

    for lot, product, routing_template in lot_rows:
        if not _is_purchase_order_target_process(
            normalized_process_type,
            routing_template.template_name,
        ):
            raise HTTPException(
                status_code=409,
                detail=(
                    f"LOT {lot.lot_no} cannot be purchase ordered for "
                    f"process {normalized_process_type}"
                ),
            )

    form_snapshot_json = None

    if payload.form_snapshot:
        if normalized_process_type == "CUT":
            form_snapshot_json = OutsourcePurchaseOrderCutSnapshot.model_validate(
                payload.form_snapshot
            ).model_dump()
        elif normalized_process_type == "PRINT":
            form_snapshot_json = OutsourcePurchaseOrderPrintSnapshot.model_validate(
                payload.form_snapshot
            ).model_dump()

    purchase_order = OutsourcePurchaseOrder(
        purchase_order_no=_generate_purchase_order_no(
            db,
            normalized_process_type,
            payload.purchase_order_date,
        ),
        purchase_order_date=payload.purchase_order_date,
        due_date=payload.due_date,
        process_type=normalized_process_type,
        outsource_partner_id=outsource_partner.partner_id,
        inbound_partner_id=payload.inbound_partner_id,
        work_description=payload.work_description,
        remark=payload.remark,
        form_snapshot_json=form_snapshot_json,
        qty=payload.qty,
        unit_price=payload.unit_price,
        supply_amount=payload.supply_amount,
        vat_amount=payload.vat_amount,
        total_amount=payload.total_amount,
    )
    db.add(purchase_order)
    db.flush()

    for item in payload.items:
        db.add(
            OutsourcePurchaseOrderItem(
                outsource_purchase_order_id=purchase_order.outsource_purchase_order_id,
                lot_id=item.lot_id,
                outsource_work_instruction_id=item.outsource_work_instruction_id,
                item_seq=item.item_seq,
                qty=item.qty,
                status=None,
            )
        )

    db.commit()
    db.refresh(purchase_order)

    return _build_purchase_order_out(db, purchase_order)


@router.get(
    "/purchase-orders/{outsource_purchase_order_id}",
    response_model=OutsourcePurchaseOrderOut,
)
def get_outsource_purchase_order(
    outsource_purchase_order_id: int,
    db: Session = Depends(get_db),
):
    purchase_order = db.get(OutsourcePurchaseOrder, outsource_purchase_order_id)

    if not purchase_order:
        raise HTTPException(status_code=404, detail="Outsource purchase order not found")

    return _build_purchase_order_out(db, purchase_order)


@router.post(
    "/purchase-orders/items/{outsource_purchase_order_item_id}/vendor-receive",
    response_model=OutsourcePurchaseOrderItemOut,
)
def vendor_receive_outsource_purchase_order_item(
    outsource_purchase_order_item_id: int,
    db: Session = Depends(get_db),
):
    item = db.get(OutsourcePurchaseOrderItem, outsource_purchase_order_item_id)

    if not item:
        raise HTTPException(status_code=404, detail="Outsource purchase order item not found")

    if item.status is not None:
        raise HTTPException(status_code=409, detail="Only not-started item can be vendor received")

    item.status = "VENDOR_RECEIVED"
    item.vendor_received_at = _utcnow()

    db.commit()
    db.refresh(item)

    lot = db.get(Lot, item.lot_id)
    order_line = db.get(OrderLine, lot.order_line_id) if lot else None
    product = db.get(Product, lot.product_id) if lot else None

    return OutsourcePurchaseOrderItemOut(
        outsource_purchase_order_item_id=item.outsource_purchase_order_item_id,
        outsource_purchase_order_id=item.outsource_purchase_order_id,
        lot_id=item.lot_id,
        outsource_work_instruction_id=item.outsource_work_instruction_id,
        item_seq=item.item_seq,
        qty=item.qty,
        status=item.status,
        vendor_received_at=item.vendor_received_at,
        work_done_at=item.work_done_at,
        shipped_at=item.shipped_at,
        work_done_qty=item.work_done_qty,
        bad_qty=item.bad_qty,
        work_done_remark=item.work_done_remark,
        created_at=item.created_at,
        lot_no=lot.lot_no if lot else None,
        order_no=order_line.order_no if order_line else None,
        line_no=order_line.line_no if order_line else None,
        product_code=product.product_code if product else None,
        product_name=product.product_name if product else None,
        lot_qty=lot.lot_qty if lot else None,
    )


@router.post(
    "/purchase-orders/items/{outsource_purchase_order_item_id}/work-done",
    response_model=OutsourcePurchaseOrderItemOut,
)
def work_done_outsource_purchase_order_item(
    outsource_purchase_order_item_id: int,
    payload: OutsourcePurchaseOrderWorkDone,
    db: Session = Depends(get_db),
):
    item = db.get(OutsourcePurchaseOrderItem, outsource_purchase_order_item_id)

    if not item:
        raise HTTPException(status_code=404, detail="Outsource purchase order item not found")

    if item.status != "VENDOR_RECEIVED":
        raise HTTPException(status_code=409, detail="Only VENDOR_RECEIVED item can be work done")

    if payload.work_done_qty + payload.bad_qty > item.qty:
        raise HTTPException(status_code=409, detail="work_done_qty + bad_qty cannot exceed qty")

    item.status = "WORK_DONE"
    item.work_done_at = _utcnow()
    item.work_done_qty = payload.work_done_qty
    item.bad_qty = payload.bad_qty
    item.work_done_remark = payload.work_done_remark

    db.commit()
    db.refresh(item)

    lot = db.get(Lot, item.lot_id)
    order_line = db.get(OrderLine, lot.order_line_id) if lot else None
    product = db.get(Product, lot.product_id) if lot else None

    return OutsourcePurchaseOrderItemOut(
        outsource_purchase_order_item_id=item.outsource_purchase_order_item_id,
        outsource_purchase_order_id=item.outsource_purchase_order_id,
        lot_id=item.lot_id,
        outsource_work_instruction_id=item.outsource_work_instruction_id,
        item_seq=item.item_seq,
        qty=item.qty,
        status=item.status,
        vendor_received_at=item.vendor_received_at,
        work_done_at=item.work_done_at,
        shipped_at=item.shipped_at,
        work_done_qty=item.work_done_qty,
        bad_qty=item.bad_qty,
        work_done_remark=item.work_done_remark,
        created_at=item.created_at,
        lot_no=lot.lot_no if lot else None,
        order_no=order_line.order_no if order_line else None,
        line_no=order_line.line_no if order_line else None,
        product_code=product.product_code if product else None,
        product_name=product.product_name if product else None,
        lot_qty=lot.lot_qty if lot else None,
    )


@router.post(
    "/purchase-orders/items/{outsource_purchase_order_item_id}/ship",
    response_model=OutsourcePurchaseOrderItemOut,
)
def ship_outsource_purchase_order_item(
    outsource_purchase_order_item_id: int,
    db: Session = Depends(get_db),
):
    item = db.get(OutsourcePurchaseOrderItem, outsource_purchase_order_item_id)

    if not item:
        raise HTTPException(status_code=404, detail="Outsource purchase order item not found")

    if item.status != "WORK_DONE":
        raise HTTPException(status_code=409, detail="Only WORK_DONE item can be shipped")

    item.status = "SHIPPED"
    item.shipped_at = _utcnow()

    db.commit()
    db.refresh(item)

    lot = db.get(Lot, item.lot_id)
    order_line = db.get(OrderLine, lot.order_line_id) if lot else None
    product = db.get(Product, lot.product_id) if lot else None

    return OutsourcePurchaseOrderItemOut(
        outsource_purchase_order_item_id=item.outsource_purchase_order_item_id,
        outsource_purchase_order_id=item.outsource_purchase_order_id,
        lot_id=item.lot_id,
        outsource_work_instruction_id=item.outsource_work_instruction_id,
        item_seq=item.item_seq,
        qty=item.qty,
        status=item.status,
        vendor_received_at=item.vendor_received_at,
        work_done_at=item.work_done_at,
        shipped_at=item.shipped_at,
        work_done_qty=item.work_done_qty,
        bad_qty=item.bad_qty,
        work_done_remark=item.work_done_remark,
        created_at=item.created_at,
        lot_no=lot.lot_no if lot else None,
        order_no=order_line.order_no if order_line else None,
        line_no=order_line.line_no if order_line else None,
        product_code=product.product_code if product else None,
        product_name=product.product_name if product else None,
        lot_qty=lot.lot_qty if lot else None,
    )


@router.get(
    "/purchase-orders",
    response_model=OutsourcePurchaseOrderListOut,
)
def get_outsource_purchase_orders(
    date_from: date | None = Query(default=None),
    date_to: date | None = Query(default=None),
    process_type: str | None = Query(default=None),
    q: str | None = Query(default=None),
    db: Session = Depends(get_db),
):
    stmt = (
        select(OutsourcePurchaseOrder)
        .order_by(
            OutsourcePurchaseOrder.purchase_order_date.desc(),
            OutsourcePurchaseOrder.outsource_purchase_order_id.desc(),
        )
    )

    if date_from:
        stmt = stmt.where(OutsourcePurchaseOrder.purchase_order_date >= date_from)

    if date_to:
        stmt = stmt.where(OutsourcePurchaseOrder.purchase_order_date <= date_to)

    normalized_process_type = (process_type or "").strip().upper()

    if normalized_process_type in {"CUT", "PRINT"}:
        stmt = stmt.where(OutsourcePurchaseOrder.process_type == normalized_process_type)

    if q and q.strip():
        like = f"%{q.strip()}%"
        stmt = stmt.where(
            (OutsourcePurchaseOrder.purchase_order_no.like(like))
            | (OutsourcePurchaseOrder.remark.like(like))
        )

    purchase_orders = db.execute(stmt).scalars().all()

    items: list[OutsourcePurchaseOrderListItemOut] = []

    for purchase_order in purchase_orders:
        outsource_partner = db.get(Partner, purchase_order.outsource_partner_id)

        items.append(
            OutsourcePurchaseOrderListItemOut(
                outsource_purchase_order_id=purchase_order.outsource_purchase_order_id,
                purchase_order_no=purchase_order.purchase_order_no,
                purchase_order_date=purchase_order.purchase_order_date,
                process_type=purchase_order.process_type,
                outsource_partner_id=purchase_order.outsource_partner_id,
                outsource_partner_name=outsource_partner.name if outsource_partner else None,
                qty=purchase_order.qty,
                remark=purchase_order.remark,
                created_at=purchase_order.created_at,
            )
        )

    return OutsourcePurchaseOrderListOut(items=items)