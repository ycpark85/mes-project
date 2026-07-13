from __future__ import annotations

from dataclasses import dataclass
from functools import lru_cache
from io import BytesIO
from pathlib import Path

from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.models.outsource_purchase_order import OutsourcePurchaseOrder
from app.schemas.outsource_work_instruction import OutsourcePurchaseOrderOut
from app.services.outsource_purchase_order_query import build_purchase_order_out


EXCEL_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@dataclass(frozen=True)
class PurchaseOrderExcelDownload:
    file_bytes: bytes
    filename: str
    media_type: str = EXCEL_MEDIA_TYPE


def build_purchase_order_excel_download(
    db: Session,
    outsource_purchase_order_id: int,
) -> PurchaseOrderExcelDownload:
    purchase_order = db.get(OutsourcePurchaseOrder, outsource_purchase_order_id)

    if not purchase_order:
        raise HTTPException(status_code=404, detail="Outsource purchase order not found")

    result = build_purchase_order_out(db, purchase_order)
    file_bytes = build_purchase_order_excel_bytes(
        result,
        purchase_order.form_snapshot_json,
    )
    filename = build_purchase_order_excel_filename(result)

    return PurchaseOrderExcelDownload(
        file_bytes=file_bytes,
        filename=filename,
    )


def build_purchase_order_excel_bytes(
    purchase_order: OutsourcePurchaseOrderOut,
    form_snapshot: dict | None,
) -> bytes:
    if purchase_order.process_type == "PRINT":
        return _build_print_purchase_order_excel_template_bytes(
            purchase_order,
            form_snapshot,
        )

    return _build_cut_purchase_order_excel_template_bytes(
        purchase_order,
        form_snapshot,
    )


def build_purchase_order_excel_filename(
    purchase_order: OutsourcePurchaseOrderOut,
) -> str:
    return f"{purchase_order.purchase_order_no}.xlsx"


@lru_cache(maxsize=1)
def _get_cut_template_bytes() -> bytes:
    template_path = (
        Path(__file__).resolve().parents[1]
        / "templates"
        / "outsource_purchase_order_cut_template.xlsx"
    )
    return template_path.read_bytes()


@lru_cache(maxsize=1)
def _get_print_template_bytes() -> bytes:
    template_path = (
        Path(__file__).resolve().parents[1]
        / "templates"
        / "outsource_purchase_order_print_template.xlsx"
    )
    return template_path.read_bytes()


def _build_print_purchase_order_excel_template_bytes(
    purchase_order: OutsourcePurchaseOrderOut,
    form_snapshot: dict | None,
) -> bytes:
    wb = load_workbook(BytesIO(_get_print_template_bytes()))
    ws = wb.active
    merged_map = _build_merged_cell_map(ws)

    snapshot = form_snapshot or {}
    rows = snapshot.get("rows") or []

    _set_merged_safe(ws, merged_map, "D5", snapshot.get("request_company_name") or "")
    _set_merged_safe(ws, merged_map, "F5", snapshot.get("requester_name") or "")
    _set_merged_safe(
        ws,
        merged_map,
        "J5",
        snapshot.get("purchase_order_date") or str(purchase_order.purchase_order_date),
    )

    start_row = 8
    max_rows = 16

    for idx, row_data in enumerate(rows[:max_rows]):
        r = start_row + idx

        _set_merged_safe(ws, merged_map, f"B{r}", row_data.get("no", ""))
        _set_merged_safe(ws, merged_map, f"C{r}", row_data.get("customer_name", ""))
        _set_merged_safe(ws, merged_map, f"D{r}", row_data.get("product_name", ""))
        _set_merged_safe(ws, merged_map, f"E{r}", row_data.get("material_spec", ""))
        _set_merged_safe(ws, merged_map, f"F{r}", row_data.get("print_sheet_qty", ""))
        _set_merged_safe(ws, merged_map, f"G{r}", row_data.get("sample", ""))
        _set_merged_safe(ws, merged_map, f"H{r}", row_data.get("plate_count", ""))
        _set_merged_safe(ws, merged_map, f"I{r}", row_data.get("color_name", ""))
        _set_merged_safe(ws, merged_map, f"J{r}", row_data.get("material_type", ""))
        _set_merged_safe(ws, merged_map, f"K{r}", row_data.get("remark", ""))

    _set_merged_safe(ws, merged_map, "D24", snapshot.get("footer_remark") or "")

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream.getvalue()


def _build_merged_cell_map(ws) -> dict[str, str]:
    merged_map: dict[str, str] = {}

    for merged_range in ws.merged_cells.ranges:
        start_ref = merged_range.start_cell.coordinate

        for row in ws[merged_range.coord]:
            for cell in row:
                merged_map[cell.coordinate] = start_ref

    return merged_map


def _set_merged_safe(ws, merged_map: dict[str, str], cell_ref: str, value) -> None:
    target_ref = merged_map.get(cell_ref, cell_ref)
    ws[target_ref] = value


def _build_cut_purchase_order_excel_template_bytes(
    purchase_order: OutsourcePurchaseOrderOut,
    form_snapshot: dict | None,
) -> bytes:
    wb = load_workbook(BytesIO(_get_cut_template_bytes()))
    ws = wb.active
    merged_map = _build_merged_cell_map(ws)

    snapshot = form_snapshot or {}
    rows = snapshot.get("rows") or []

    _set_merged_safe(ws, merged_map, "E5", snapshot.get("request_company_name") or "")
    _set_merged_safe(ws, merged_map, "I5", snapshot.get("requester_name") or "")
    _set_merged_safe(
        ws,
        merged_map,
        "P5",
        snapshot.get("purchase_order_date") or str(purchase_order.purchase_order_date),
    )
    _set_merged_safe(
        ws,
        merged_map,
        "H7",
        snapshot.get("raw_material_inbound_text") or "",
    )

    start_row = 10
    max_rows = 16

    for idx, row_data in enumerate(rows[:max_rows]):
        r = start_row + idx

        _set_merged_safe(ws, merged_map, f"B{r}", row_data.get("no", ""))
        _set_merged_safe(ws, merged_map, f"C{r}", row_data.get("raw_material_text", ""))
        _set_merged_safe(ws, merged_map, f"H{r}", row_data.get("length_m_text", ""))
        _set_merged_safe(ws, merged_map, f"I{r}", row_data.get("inbound_place_text", ""))
        _set_merged_safe(ws, merged_map, f"O{r}", row_data.get("cut_spec_text", ""))
        _set_merged_safe(ws, merged_map, f"P{r}", row_data.get("sheet_qty_text", ""))

    _set_merged_safe(ws, merged_map, "H30", snapshot.get("stock_500_width_text") or "")
    _set_merged_safe(ws, merged_map, "J30", snapshot.get("stock_600_width_text") or "")
    _set_merged_safe(ws, merged_map, "O30", snapshot.get("stock_600_tpt0268_text") or "")
    _set_merged_safe(ws, merged_map, "B32", snapshot.get("remark") or "")

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream.getvalue()
