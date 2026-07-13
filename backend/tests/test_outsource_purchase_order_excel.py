from __future__ import annotations

import unittest
from datetime import date, datetime
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import Mock, patch

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook

import app.services.outsource_purchase_order_excel as excel
from app.schemas.outsource_work_instruction import OutsourcePurchaseOrderOut


class OutsourcePurchaseOrderExcelTests(unittest.TestCase):
    def test_build_cut_purchase_order_excel_bytes_fills_snapshot_cells(self) -> None:
        purchase_order = _build_purchase_order(process_type="CUT")
        snapshot = {
            "request_company_name": "Request Co",
            "requester_name": "Kim",
            "purchase_order_date": "2026-01-03",
            "raw_material_inbound_text": "Inbound",
            "stock_500_width_text": "500 stock",
            "stock_600_width_text": "600 stock",
            "stock_600_tpt0268_text": "TPT stock",
            "remark": "Cut remark",
            "rows": [
                {
                    "no": 1,
                    "raw_material_text": "Material",
                    "length_m_text": "100m",
                    "inbound_place_text": "Bohyun",
                    "cut_spec_text": "10x20",
                    "sheet_qty_text": "30",
                }
            ],
        }

        with patch.object(excel, "_get_cut_template_bytes", return_value=_blank_workbook_bytes()):
            result = excel.build_purchase_order_excel_bytes(purchase_order, snapshot)

        ws = _load_active_sheet(result)

        self.assertEqual("Request Co", ws["E5"].value)
        self.assertEqual("Kim", ws["I5"].value)
        self.assertEqual("2026-01-03", ws["P5"].value)
        self.assertEqual("Inbound", ws["H7"].value)
        self.assertEqual("Material", ws["C10"].value)
        self.assertEqual("100m", ws["H10"].value)
        self.assertEqual("Cut remark", ws["B32"].value)

    def test_build_print_purchase_order_excel_bytes_fills_snapshot_cells(self) -> None:
        purchase_order = _build_purchase_order(process_type="PRINT")
        snapshot = {
            "request_company_name": "Request Co",
            "requester_name": "Lee",
            "purchase_order_date": "2026-01-04",
            "footer_remark": "Print remark",
            "rows": [
                {
                    "no": 1,
                    "customer_name": "Customer",
                    "product_name": "Product",
                    "material_spec": "Spec",
                    "print_sheet_qty": "40",
                    "sample": "Y",
                    "plate_count": "2",
                    "color_name": "Black",
                    "material_type": "PET",
                    "remark": "Row remark",
                }
            ],
        }

        with patch.object(excel, "_get_print_template_bytes", return_value=_blank_workbook_bytes()):
            result = excel.build_purchase_order_excel_bytes(purchase_order, snapshot)

        ws = _load_active_sheet(result)

        self.assertEqual("Request Co", ws["D5"].value)
        self.assertEqual("Lee", ws["F5"].value)
        self.assertEqual("2026-01-04", ws["J5"].value)
        self.assertEqual("Customer", ws["C8"].value)
        self.assertEqual("Product", ws["D8"].value)
        self.assertEqual("Row remark", ws["K8"].value)
        self.assertEqual("Print remark", ws["D24"].value)

    def test_build_purchase_order_excel_filename_uses_purchase_order_no(self) -> None:
        purchase_order = _build_purchase_order(process_type="CUT")

        self.assertEqual(
            "OCUT-20260103-001.xlsx",
            excel.build_purchase_order_excel_filename(purchase_order),
        )

    def test_build_purchase_order_excel_download_returns_download_payload(self) -> None:
        db = Mock()
        db_purchase_order = SimpleNamespace(
            outsource_purchase_order_id=1,
            form_snapshot_json={"rows": []},
        )
        purchase_order = _build_purchase_order(process_type="CUT")
        db.get.return_value = db_purchase_order

        with (
            patch.object(excel, "build_purchase_order_out", return_value=purchase_order),
            patch.object(excel, "build_purchase_order_excel_bytes", return_value=b"xlsx") as bytes_mock,
        ):
            result = excel.build_purchase_order_excel_download(db, 1)

        db.get.assert_called_once()
        bytes_mock.assert_called_once_with(purchase_order, {"rows": []})
        self.assertEqual(b"xlsx", result.file_bytes)
        self.assertEqual("OCUT-20260103-001.xlsx", result.filename)
        self.assertEqual(excel.EXCEL_MEDIA_TYPE, result.media_type)

    def test_build_purchase_order_excel_download_rejects_missing_order(self) -> None:
        db = Mock()
        db.get.return_value = None

        with self.assertRaises(HTTPException) as ctx:
            excel.build_purchase_order_excel_download(db, 999)

        self.assertEqual(404, ctx.exception.status_code)
        self.assertEqual("Outsource purchase order not found", ctx.exception.detail)


def _blank_workbook_bytes() -> bytes:
    stream = BytesIO()
    Workbook().save(stream)
    stream.seek(0)
    return stream.getvalue()


def _load_active_sheet(file_bytes: bytes):
    return load_workbook(BytesIO(file_bytes)).active


def _build_purchase_order(process_type: str) -> OutsourcePurchaseOrderOut:
    return OutsourcePurchaseOrderOut(
        outsource_purchase_order_id=1,
        purchase_order_no="OCUT-20260103-001",
        purchase_order_date=date(2026, 1, 3),
        due_date=None,
        process_type=process_type,
        outsource_partner_id=1,
        inbound_partner_id=None,
        work_description=None,
        remark=None,
        qty=1,
        created_at=datetime(2026, 1, 3, 9, 0, 0),
        updated_at=datetime(2026, 1, 3, 9, 0, 0),
        items=[],
    )


if __name__ == "__main__":
    unittest.main()
